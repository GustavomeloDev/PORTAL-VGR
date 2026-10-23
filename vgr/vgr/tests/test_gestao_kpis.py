"""Testes da Gestão de Atendimentos VGR (KPIs do período + relatório)."""
from __future__ import annotations

import io
import unittest
from datetime import date, datetime, timedelta

from vgr.vgr.gestao.gestao_kpis import (
    ESCOPO_AGENDADO_NO_PERIODO,
    ESCOPO_AGENDADO_PENDENTE,
    ESCOPO_FORA_DE_ESCOPO,
    ESCOPO_IMPRODUTIVA,
    calcular_kpis,
    escopo_linha,
    fmt_hhmmss,
    mesclar_linhas,
    parse_periodo,
    resumo_periodo_payload,
)
from vgr.vgr.gestao.relatorio_atividades import (
    COLUNAS_ATIVIDADES,
    gerar_csv,
    gerar_xlsx,
    linha_atividade,
)

HOJE = date(2026, 10, 13)


def _ag(id_ag, data, status_ag="Agendado", **extra):
    row = {
        "id_agendamento": id_ag,
        "data_agendamento": data,
        "hora_agendamento": timedelta(hours=9, minutes=30),
        "status_agendamento": status_ag,
        "id_atendimento": None,
        "razao_social": "Cliente X",
        "nome_unidade": f"Unidade {id_ag}",
    }
    row.update(extra)
    return row


def _at(id_at, status, sub, *, data_ag=None, abertura=None, fechamento=None, **extra):
    row = {
        "id_agendamento": 900 + id_at if data_ag else None,
        "data_agendamento": data_ag,
        "id_atendimento": id_at,
        "status_atendimento": status,
        "substatus_atendimento": sub,
        "abertura": abertura,
        "fechamento": fechamento,
        "data_atividade": abertura,
        "razao_social": "Cliente X",
        "nome_unidade": f"Unidade A{id_at}",
    }
    row.update(extra)
    return row


class TestPeriodo(unittest.TestCase):
    def test_padrao_hoje(self):
        self.assertEqual(parse_periodo(None, None, hoje=HOJE), (HOJE, HOJE))

    def test_intervalo_livre(self):
        s, e = parse_periodo("2026-10-07", "2026-10-13", hoje=HOJE)
        self.assertEqual((s, e), (date(2026, 10, 7), HOJE))

    def test_inicio_maior_que_fim(self):
        with self.assertRaises(ValueError):
            parse_periodo("2026-10-14", "2026-10-13", hoje=HOJE)

    def test_formato_invalido(self):
        with self.assertRaises(ValueError):
            parse_periodo("13/10/2026", "2026-10-13", hoje=HOJE)

    def test_limite_maximo(self):
        with self.assertRaises(ValueError):
            parse_periodo("2025-01-01", "2026-10-13", hoje=HOJE)


class TestKpis(unittest.TestCase):
    def setUp(self):
        ab = datetime(2026, 10, 13, 9, 0)
        self.rows = [
            _ag(1, HOJE),  # agendado pendente
            _ag(2, HOJE, status_ag="Improdutiva"),  # fora do total
            _at(10, "Agendado", "Check-list", data_ag=HOJE),  # preparação = agendado
            _at(11, "Em Atendimento", "Aguarda Gestão", data_ag=HOJE),
            _at(12, "Em Atendimento", "Aguarda Cliente", data_ag=HOJE, abertura=ab),
            _at(13, "Em Atendimento", "Aguarda Contratada", abertura=ab),  # walk-in
            _at(14, "Em Atendimento", "Em Configuração Remota", data_ag=HOJE, abertura=ab),
            _at(15, "Concluído", "Concluído", data_ag=HOJE, abertura=ab,
                fechamento=ab + timedelta(hours=2)),
            _at(16, "Concluído", "Day After", data_ag=date(2026, 10, 10), abertura=ab,
                fechamento=ab + timedelta(hours=4)),
            _at(17, "Quebra", "Quebra Field", data_ag=HOJE, abertura=ab,
                fechamento=ab + timedelta(hours=1)),
            _at(18, "Concluído Parcial", "Pendente Link IPD", abertura=ab,
                fechamento=ab + timedelta(hours=1)),
        ]

    def test_contagens(self):
        k = calcular_kpis(self.rows, HOJE, HOJE)["kpis"]
        self.assertEqual(k["agendados"], 2)
        self.assertEqual(k["com_atividade"], 8)
        self.assertEqual(k["total"], 10)
        self.assertEqual(k["improdutivas"], 1)
        self.assertEqual(k["em_atendimento"], 4)
        self.assertEqual(k["aguarda_gestao"], 1)
        self.assertEqual(k["aguarda_cliente"], 1)
        self.assertEqual(k["aguarda_contratada"], 1)
        self.assertEqual(k["fluxo_tecnico"], 1)
        self.assertEqual(k["concluidos"], 1)
        self.assertEqual(k["day_after"], 1)
        self.assertEqual(k["quebra"], 1)
        self.assertEqual(k["concluido_parcial"], 1)
        self.assertEqual(k["pct_concluidos"], 10)

    def test_tempo_medio_so_finalizados(self):
        k = calcular_kpis(self.rows, HOJE, HOJE)["kpis"]
        # 2h + 4h + 1h + 1h = 8h / 4
        self.assertEqual(k["tempo_medio"], "02:00:00")
        self.assertEqual(k["tempo_medio_amostras"], 4)

    def test_detalhe_total_bate_com_total(self):
        calc = calcular_kpis(self.rows, HOJE, HOJE)
        det, k = calc["detalhe_total"], calc["kpis"]
        self.assertEqual(det["total"], k["total"])
        self.assertEqual(det["agendados_pend"], 2)
        self.assertEqual(det["agendados_ativos"], 3)  # 11, 12, 14
        self.assertEqual(det["agendados_finalizados"], 2)  # 15, 17
        self.assertEqual(det["fora_escopo_ativos"], 1)  # 13 walk-in
        self.assertEqual(det["fora_escopo_finalizados"], 2)  # 16 agenda antiga, 18 walk-in

    def test_payload_hint_e_periodo(self):
        p = resumo_periodo_payload(self.rows, HOJE, HOJE)
        self.assertIn("2 Agendados Pend.", p["hint_total"])
        self.assertIn("1 improdutiva(s)", p["hint_total"])
        self.assertEqual(p["periodo"]["dias"], 1)

    def test_vazio(self):
        k = calcular_kpis([], HOJE, HOJE)["kpis"]
        self.assertEqual(k["total"], 0)
        self.assertEqual(k["pct_concluidos"], 0)
        self.assertEqual(k["tempo_medio"], "00:00:00")


class TestEscopoEMescla(unittest.TestCase):
    def test_escopos(self):
        self.assertEqual(escopo_linha(_ag(1, HOJE), HOJE, HOJE), ESCOPO_AGENDADO_PENDENTE)
        self.assertEqual(
            escopo_linha(_ag(2, HOJE, status_ag="Improdutiva"), HOJE, HOJE),
            ESCOPO_IMPRODUTIVA,
        )
        self.assertEqual(
            escopo_linha(_at(3, "Quebra", "Quebra Link", data_ag=HOJE), HOJE, HOJE),
            ESCOPO_AGENDADO_NO_PERIODO,
        )
        self.assertEqual(
            escopo_linha(_at(4, "Quebra", "Quebra Link"), HOJE, HOJE),
            ESCOPO_FORA_DE_ESCOPO,
        )

    def test_mescla_sem_duplicar_atendimento(self):
        a = [_at(10, "Em Atendimento", "Aguarda Gestão", data_ag=HOJE), _ag(1, HOJE)]
        b = [_at(10, "Em Atendimento", "Aguarda Gestão", data_ag=HOJE), _at(11, "Quebra", "x")]
        self.assertEqual(len(mesclar_linhas(a, b)), 3)

    def test_fmt(self):
        self.assertEqual(fmt_hhmmss(3661), "01:01:01")
        self.assertEqual(fmt_hhmmss(None), "00:00:00")
        self.assertEqual(fmt_hhmmss(-5), "00:00:00")


class TestRelatorio(unittest.TestCase):
    def test_linha_tem_todas_as_colunas_e_bloqueia_formula(self):
        row = _ag(1, HOJE, razao_social="=HYPERLINK(\"x\")")
        linha = linha_atividade(row, HOJE, HOJE)
        self.assertEqual(len(linha), len(COLUNAS_ATIVIDADES))
        self.assertTrue(linha[2].startswith("'="))
        self.assertEqual(linha[9], "13/10/2026")
        self.assertEqual(linha[10], "09:30")
        self.assertEqual(linha[16], "Agendado")

    def test_csv_bom_e_separador(self):
        data = gerar_csv([_ag(1, HOJE)], HOJE, HOJE)
        self.assertTrue(data.startswith("﻿".encode("utf-8")))
        self.assertIn(b"Escopo;ID Atendimento;", data)

    def test_xlsx_abas(self):
        from openpyxl import load_workbook

        rows = [_ag(1, HOJE), _at(5, "Concluído", "Concluído", data_ag=HOJE,
                                  abertura=datetime(2026, 10, 13, 8),
                                  fechamento=datetime(2026, 10, 13, 9))]
        resumo = resumo_periodo_payload(rows, HOJE, HOJE)
        trans = [{
            "criado_em": datetime(2026, 10, 13, 9), "id_atendimento": 5,
            "razao_social": "Cliente X", "nome_unidade": "U",
            "status_anterior": "Em Atendimento", "substatus_anterior": "Em Caderno de Teste",
            "status_novo": "Concluído", "substatus_novo": "Concluído",
            "usuario_re": "123", "usuario_nome": "Ana", "observacao": None,
        }]
        data = gerar_xlsx(rows, trans, resumo, HOJE, HOJE)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Resumo", "Atividades", "Transições"])
        ws = wb["Atividades"]
        self.assertEqual(ws.cell(row=2, column=1).value, "Escopo")
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(row=4, column=20).value, "01:00:00")
        self.assertEqual(wb["Transições"].cell(row=3, column=9).value, "123 – Ana")


if __name__ == "__main__":
    unittest.main()
