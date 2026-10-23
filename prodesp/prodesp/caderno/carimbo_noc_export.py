"""
Texto plano da aba CARIMBO NOC (coluna E) para copiar ao clipboard.
"""
from __future__ import annotations

import glob
import os
import re
import warnings
from typing import Any, Optional

from prodesp.prodesp.caderno.caderno_carimbo_links import deve_exibir_bloco_sai_bli_no_carimbo
from prodesp.prodesp.caderno.caderno_template_map import CARIMBO_CMD_ANCHOR_ROWS, CARIMBO_LINK_SAI_BLI
from prodesp.prodesp.caderno.gerador_caderno_prodesp import resolve_gerados_dir, _parse_logs_por_comando

_CARIMBO_E_START_ROW = 7
_CARIMBO_E_COL = 5
_SHOW_IFACE_PREFIX = "show sys interface "
_CMD_SHOW_FULL = "show sys interface"
_SAI_BLI_ROW_START = min(CARIMBO_LINK_SAI_BLI.values()) - 2
_SAI_BLI_ROW_END = max(CARIMBO_LINK_SAI_BLI.values())


def _linhas_bloco_sai_bli() -> frozenset[int]:
    return frozenset(range(_SAI_BLI_ROW_START, _SAI_BLI_ROW_END + 1))


def _ler_bloco_sai_bli_coluna_d(ws) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for field, row in CARIMBO_LINK_SAI_BLI.items():
        val = ws.cell(row, 4).value
        if val is None:
            continue
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        out[field] = val
    return out


def _remover_bloco_sai_bli_vazio_texto(texto: str) -> str:
    """Remove seção Conectividade SAI-BLI quando circuito/id estão N/A ou vazios."""
    if not texto or "Conectividade SAI-BLI" not in texto:
        return texto
    partes = re.split(r"(?=Conectividade )", texto)
    filtradas: list[str] = []
    for parte in partes:
        if not parte.strip():
            continue
        if parte.strip().startswith("Conectividade SAI-BLI"):
            linhas = parte.strip().splitlines()
            campos: dict[str, str] = {}
            for linha in linhas[1:]:
                if ":" not in linha:
                    continue
                chave, _, valor = linha.partition(":")
                campos[chave.strip().lower().replace(" ", "_")] = valor.strip()
            if not deve_exibir_bloco_sai_bli_no_carimbo(
                {
                    "circuito": campos.get("circuito"),
                    "id_vantive": campos.get("id_vantive"),
                    "velocidade": campos.get("velocidade"),
                    "description": campos.get("description"),
                    "ip_wan": campos.get("ip_wan"),
                }
            ):
                continue
        filtradas.append(parte.strip())
    return _compactar_linhas_em_branco("\n\n".join(filtradas))


def _compactar_linhas_em_branco(texto: str) -> str:
    """Mantém no máximo uma linha em branco consecutiva entre blocos."""
    if not texto:
        return ""
    linhas = texto.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    saida: list[str] = []
    em_branco = False
    for linha in linhas:
        if not linha.strip():
            if not em_branco and saida:
                saida.append("")
                em_branco = True
            continue
        em_branco = False
        saida.append(linha.rstrip())
    return "\n".join(saida).strip()


def _juntar_blocos(blocos: list[str]) -> str:
    partes = [b.strip() for b in blocos if (b or "").strip()]
    return _compactar_linhas_em_branco("\n\n".join(partes))


def _is_linha_comando(texto: str, comandos: set[str]) -> bool:
    t = (texto or "").strip()
    if not t:
        return False
    if t in comandos:
        return True
    if t == _CMD_SHOW_FULL:
        return True
    return t.startswith(_SHOW_IFACE_PREFIX)


def _dirs_caderno_busca(acronimo: str = "") -> list[str]:
    dirs = [resolve_gerados_dir()]
    try:
        from prodesp.config_paths import iter_pastas_rede_caderno

        for rede in iter_pastas_rede_caderno((acronimo or "").strip()):
            if rede and rede not in dirs:
                dirs.append(rede)
    except Exception:
        pass
    return dirs


def _ultimo_caderno_gerado(id_vgr: str, acronimo: str = "", id_intragov: str = "") -> Optional[str]:
    id_vgr = (id_vgr or "").strip()
    if not id_vgr:
        return None
    acr = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", (acronimo or "").strip())
    intra = (id_intragov or "").strip().replace("/", "-")
    matches: list[str] = []
    for base_dir in _dirs_caderno_busca(acronimo):
        patterns: list[str] = []
        for ext in (".xlsx", ".xlsm"):
            if acr and intra:
                patterns.append(os.path.join(base_dir, f"{acr}_ID_VGR_{id_vgr}_ID_INTRAGOV_{intra}{ext}"))
            if acr:
                patterns.append(os.path.join(base_dir, f"{acr}_ID_VGR_{id_vgr}_ID_INTRAGOV_*{ext}"))
            patterns.append(os.path.join(base_dir, f"*_ID_VGR_{id_vgr}_ID_INTRAGOV_*{ext}"))
            patterns.append(
                os.path.join(
                    base_dir,
                    f"Caderno_TESTE_{id_vgr}_{acr}_*{ext}" if acr else f"Caderno_TESTE_{id_vgr}_*{ext}",
                )
            )
        for pattern in patterns:
            matches.extend(glob.glob(pattern))
    matches = list(dict.fromkeys(matches))
    if not matches:
        return None
    return max(matches, key=os.path.getmtime)


def extrair_texto_carimbo_coluna_e(caminho_xlsm: str, *, start_row: int = _CARIMBO_E_START_ROW) -> str:
    """Lê coluna E (valores calculados) — um bloco por comando, separados por linha em branco."""
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl(\..*)?")
        wb = load_workbook(caminho_xlsm, read_only=True, data_only=True, keep_vba=False)
        try:
            ws = wb["CARIMBO NOC"]
            incluir_sai_bli = deve_exibir_bloco_sai_bli_no_carimbo(_ler_bloco_sai_bli_coluna_d(ws))
            linhas_omitir = set() if incluir_sai_bli else _linhas_bloco_sai_bli()

            comandos = set(CARIMBO_CMD_ANCHOR_ROWS.keys())
            celulas: list[str] = []
            max_row = ws.max_row or start_row
            for row_num, row_cells in enumerate(
                ws.iter_rows(min_row=start_row, max_row=max_row, min_col=_CARIMBO_E_COL, max_col=_CARIMBO_E_COL),
                start=start_row,
            ):
                if row_num in linhas_omitir:
                    continue
                cell = row_cells[0]
                if cell.value is None:
                    continue
                texto = str(cell.value).strip()
                if texto:
                    celulas.append(texto)

            blocos: list[str] = []
            bloco_atual: list[str] = []
            for texto in celulas:
                if _is_linha_comando(texto, comandos):
                    if bloco_atual:
                        blocos.append("\n".join(bloco_atual))
                    bloco_atual = [texto]
                else:
                    bloco_atual.append(texto)
            if bloco_atual:
                blocos.append("\n".join(bloco_atual))

            return _remover_bloco_sai_bli_vazio_texto(_juntar_blocos(blocos))
        finally:
            wb.close()


def montar_texto_carimbo_from_logs(logs_text: str) -> str:
    """Monta texto equivalente à coluna E a partir dos logs coletados (sem cabeçalho Excel)."""
    por_cmd = _parse_logs_por_comando(logs_text)
    if not por_cmd:
        return ""

    blocos: list[str] = []
    for cmd, anchor in sorted(CARIMBO_CMD_ANCHOR_ROWS.items(), key=lambda item: item[1]):
        if cmd == _CMD_SHOW_FULL or cmd.startswith(_SHOW_IFACE_PREFIX):
            continue
        body = (por_cmd.get(cmd) or "").strip()
        if body:
            body = _compactar_linhas_em_branco(body)
            blocos.append(f"{cmd}\n{body}")

    show_parts = [
        f"{cmd}\n{_compactar_linhas_em_branco(body.strip())}"
        for cmd, body in por_cmd.items()
        if (cmd == _CMD_SHOW_FULL or cmd.startswith(_SHOW_IFACE_PREFIX)) and (body or "").strip()
    ]
    if show_parts:
        blocos.append("\n".join(show_parts))

    return _juntar_blocos(blocos)


def _format_campo_carimbo(rotulo: str, valor: Any) -> str:
    v = (str(valor) if valor is not None else "").strip()
    if not v:
        v = "N/A"
    return f"{rotulo}:\t{v}"


def _format_bloco_conectividade(titulo: str, bloco: dict[str, Any]) -> str:
    if not bloco or not (bloco.get("circuito") or bloco.get("interface")):
        return ""
    rotulos = (
        ("INTERFACE", "interface"),
        ("CIRCUITO", "circuito"),
        ("TECNOLOGIA", "tecnologia"),
        ("ID VANTIVE", "id_vantive"),
        ("VELOCIDADE", "velocidade"),
        ("DESCRIPTION", "description"),
        ("IP WAN", "ip_wan"),
        ("REMOTO", "remoto"),
    )
    linhas = [f"Conectividade {titulo}"]
    for rotulo, chave in rotulos:
        val = bloco.get(chave)
        if val is None or str(val).strip() == "":
            continue
        linhas.append(_format_campo_carimbo(rotulo, val))
    return "\n".join(linhas).strip()


def montar_texto_carimbo_completo(
    *,
    unidade: dict[str, Any],
    equipamentos: list[dict[str, Any]],
    links: list[dict[str, Any]],
    questionario: dict[str, Any] | None = None,
    logs_text: str = "",
) -> str:
    """Monta carimbo completo (cabeçalho + links + logs) sem depender do Excel."""
    from prodesp.prodesp.caderno.caderno_carimbo_links import (
        deve_exibir_bloco_sai_bli_no_carimbo,
        montar_bloco_carimbo,
        resolver_links_carimbo,
    )
    from prodesp.prodesp.caderno.caderno_interface_resolver import (
        extrair_aliases_show_interface,
        extrair_arp_por_interface,
        extrair_ips_interfaces_logs,
        montar_show_interface_parts_para_carimbo,
        resolver_interfaces_carimbo,
    )
    from prodesp.prodesp.caderno.caderno_template_map import CADERNO_MODELO_SOLUCAO_VALOR
    from prodesp.prodesp.caderno.fgt_model_interfaces import normalizar_modelo_exibicao
    from prodesp.prodesp.caderno.fgt_ssh_collector import extrair_bloco_coleta_fgt

    q = questionario or {}
    eq_prim = equipamentos[0] if equipamentos else {}
    modelo_raw = eq_prim.get("modelo_equipamentos") or eq_prim.get("modelo") or ""
    modelo = normalizar_modelo_exibicao(modelo_raw)
    serials = [
        (e.get("serial_equipamentos") or e.get("serial") or "").strip()
        for e in equipamentos
        if (e.get("serial_equipamentos") or e.get("serial") or "").strip()
    ]
    serial_txt = serials[0] if len(serials) == 1 else " / ".join(serials) if serials else ""

    id_intragov = (
        (unidade.get("id_intragov_unidades") or unidade.get("id_intragov") or "")
        .replace("/", "-")
        .strip()
    )
    cabecalho = [
        "CARIMBO - SDWAN",
        "",
        _format_campo_carimbo("ID SDWAN-INTRAGOV", id_intragov),
        _format_campo_carimbo("ID SDWAN-VTV", unidade.get("id_vgr_unidades") or unidade.get("id_vgr")),
        _format_campo_carimbo("ACRÔNIMO", unidade.get("acronimo_unidades") or unidade.get("acronimo")),
        _format_campo_carimbo("SITE NAME", unidade.get("nome_unidades") or unidade.get("site_name")),
        _format_campo_carimbo("LOCALIDADE", unidade.get("localidade_unidades") or unidade.get("localidade")),
        _format_campo_carimbo("TOPOLOGIA", unidade.get("topologia_unidades") or unidade.get("topologia") or "SD-WAN"),
        _format_campo_carimbo("MODELO DE SOLUÇÃO", unidade.get("modelo_solucao_unidades") or CADERNO_MODELO_SOLUCAO_VALOR),
        _format_campo_carimbo("FORNECEDOR", "Fortinet"),
        _format_campo_carimbo("MODELO DE EQUIPAMENTO", modelo),
        _format_campo_carimbo("REDUNDANCIA (HA)", "SIM" if len(serials) > 1 else "NÃO"),
        _format_campo_carimbo("HOSTNAME PRIMARIO", (unidade.get("hostname_unidades") or unidade.get("hostname") or "").lower()),
        _format_campo_carimbo("SERIAL NUMBER", serial_txt),
        _format_campo_carimbo("IP-MGMT", unidade.get("loopback_vivo_unidades") or unidade.get("loopback_vivo")),
        _format_campo_carimbo("POSSUI DHCP-SERVER", q.get("dhcp_server") or "NÃO"),
        _format_campo_carimbo("POSSUI DHCP-RELAY", q.get("dhcp_relay") or "NÃO"),
        _format_campo_carimbo("POSSUI VLAN", q.get("vlan") or "NÃO"),
        _format_campo_carimbo("POSSUI MULTIPLE LAN IP", q.get("multi_ip_lan") or "NÃO"),
    ]

    logs_fgt = extrair_bloco_coleta_fgt(logs_text or "")
    por_cmd = _parse_logs_por_comando(logs_fgt)
    links_res = resolver_links_carimbo(links)
    aliases_iface = extrair_aliases_show_interface(por_cmd)
    iface_ips = extrair_ips_interfaces_logs(por_cmd)
    arp_por_iface = extrair_arp_por_interface(por_cmd)
    ifaces_por_link = resolver_interfaces_carimbo(aliases_iface, links_res, modelo, por_cmd)

    def _ip_log_iface(nome_iface: str) -> str:
        if not nome_iface:
            return ""
        return iface_ips.get(nome_iface.lower(), "") or ""

    def _arp_iface(nome_iface: str) -> list[str]:
        if not nome_iface:
            return []
        return arp_por_iface.get(nome_iface.lower(), []) or []

    link_sai_bl = montar_bloco_carimbo(
        links_res.sai_bl,
        interface=ifaces_por_link.get("sai_bl", ""),
        circuito="SAI-BL",
        tecnologia="INET1",
        ip_interface_log=_ip_log_iface(ifaces_por_link.get("sai_bl", "")),
        arp_neighbors=_arp_iface(ifaces_por_link.get("sai_bl", "")),
    )
    link_scm = montar_bloco_carimbo(
        links_res.scm,
        interface=ifaces_por_link.get("scm", ""),
        circuito="SCM",
        tecnologia="MPLS1",
        ip_interface_log=_ip_log_iface(ifaces_por_link.get("scm", "")),
        arp_neighbors=_arp_iface(ifaces_por_link.get("scm", "")),
    )
    link_sai_bli: dict[str, str] | None = None
    if links_res.incluir_sai_bli and ifaces_por_link.get("sai_bli"):
        candidato = montar_bloco_carimbo(
            links_res.sai_bli,
            interface=ifaces_por_link["sai_bli"],
            circuito="SAI-BLI",
            tecnologia="INET2",
            ip_interface_log=_ip_log_iface(ifaces_por_link["sai_bli"]),
            arp_neighbors=_arp_iface(ifaces_por_link["sai_bli"]),
        )
        if deve_exibir_bloco_sai_bli_no_carimbo(candidato):
            link_sai_bli = candidato

    blocos_link: list[str] = []
    for titulo, bloco in (
        (link_sai_bl.get("circuito") or "SAI-BL", link_sai_bl),
        (link_scm.get("circuito") or "SCM", link_scm),
        ((link_sai_bli or {}).get("circuito") or "SAI-BLI", link_sai_bli),
    ):
        if not bloco:
            continue
        fmt = _format_bloco_conectividade(titulo, bloco)
        if fmt:
            blocos_link.append(fmt)

    logs_part = montar_texto_carimbo_from_logs(logs_fgt) if logs_fgt else ""
    partes = ["\n".join(cabecalho).strip()]
    partes.extend(blocos_link)
    if logs_part:
        partes.append(logs_part)
    return _remover_bloco_sai_bli_vazio_texto(_juntar_blocos(partes))


def obter_texto_carimbo_noc(
    *,
    id_vgr: str,
    acronimo: str = "",
    id_intragov: str = "",
    logs_text: str = "",
    caminho_xlsm: Optional[str] = None,
    unidade: Optional[dict[str, Any]] = None,
    equipamentos: Optional[list[dict[str, Any]]] = None,
    links: Optional[list[dict[str, Any]]] = None,
    questionario: Optional[dict[str, Any]] = None,
) -> tuple[str, str]:
    """
    Retorna (texto, fonte) — fonte: 'xlsm', 'montado' ou 'logs'.
    Preferência: caderno .xlsm gerado; fallback: montagem completa; por último só logs.
    """
    path = caminho_xlsm or _ultimo_caderno_gerado(id_vgr, acronimo, id_intragov)
    if path and os.path.isfile(path):
        try:
            texto = extrair_texto_carimbo_coluna_e(path)
            if texto:
                return texto, "xlsm"
        except OSError:
            pass
        except Exception:
            pass

    if unidade and links is not None:
        try:
            texto = montar_texto_carimbo_completo(
                unidade=unidade,
                equipamentos=equipamentos or [],
                links=links,
                questionario=questionario or {},
                logs_text=logs_text or "",
            )
            if texto:
                return texto, "montado"
        except Exception:
            pass

    from prodesp.prodesp.caderno.fgt_ssh_collector import extrair_bloco_coleta_fgt

    logs_fgt = extrair_bloco_coleta_fgt(logs_text or "")
    if logs_fgt:
        texto = montar_texto_carimbo_from_logs(logs_fgt)
        if texto:
            return texto, "logs"

    return "", ""
